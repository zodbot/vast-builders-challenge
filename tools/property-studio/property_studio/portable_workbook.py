"""VM fallback, used only when the artifact-tool runtime is unavailable.

Formulas are identical to the verified shared graph. The Python graph supplies
initial caches; Excel recalculates on open. This is not native Excel QA.
"""
import io
import json
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.workbook.properties import CalcProperties


def safe(value):
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value


def export_portable(spec,path):
    wb=Workbook();wb.remove(wb.active)
    names=['Summary','Inputs','Rental Operations','Amortization','Rent vs Buy','Sensitivity','Sales Comps','Rental Comps','Property Evidence','HOA Questions','Sources']
    for name in names:
        ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
        ws.column_dimensions['A'].width=3;ws.column_dimensions['B'].width=3
        ws.column_dimensions['C'].width=48;ws.column_dimensions['D'].width=23
        ws.column_dimensions['E'].width=3;ws.column_dimensions['F'].width=80
        ws['C2']=name;ws['C2'].font=Font(name='Arial',size=16,bold=True,color='173F38')
        ws['C3']=safe(spec['address']);ws['F3']=f"{spec['mode'].upper()} · Version {spec['version']}"
        ws.freeze_panes='D6' if name!='Summary' else None
    caches={}
    for c in spec['finance']['cells']:
        ws=wb[c['sheet']];cell=ws[c['cell']]
        cell.value=c['formula'] if c.get('formula') else c['value']
        if c.get('label_cell') and c.get('label'): ws[c['label_cell']]=safe(c['label'])
        if c.get('note'):
            ws.cell(cell.row,6,safe(c['note'])).alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[cell.row].height=32
        cell.number_format='0.0%;(0.0%);"–"' if c['format']=='percent' else '0.00"x"' if c['format']=='ratio' else '$#,##0.00;($#,##0.00);"–"' if c['format']=='money' else '#,##0.00'
        cell.font=Font(name='Arial',size=10,color='1764A1' if c.get('input') else '173F38')
        if c.get('input'): cell.fill=PatternFill('solid',fgColor='EDF4F1' if c['value'] is not None else 'FFF1C9')
        if c.get('formula'): caches[(c['sheet'],c['cell'])]=c['value']
    for name in ['Amortization','Rent vs Buy']:
        ws=wb[name]
        headers=['Month','Payment','Interest','Principal','Balance'] if name=='Amortization' else ['Month','Owner cash cost','Renter cash cost','Renter invested','Buyer invested']
        for i,h in enumerate(headers,3): ws.cell(5,i,h);ws.column_dimensions[chr(64+i)].width=23
        for row in range(6,ws.max_row+1): ws.cell(row,3,row-5)
    for i,label in enumerate(['Hypothetical case','Purchase price','Monthly rent','Monthly P&I','Annual NOI','Cap rate','Annual cash flow'],3):
        wb['Sensitivity'].cell(5,i,label)
        wb['Sensitivity'].column_dimensions[chr(64+i)].width=25
    wb['Sensitivity']['C17']='Tax budget scales with purchase price. Other costs stay fixed. Not market forecasts.'
    def table(name,headers,rows):
        ws=wb[name]
        for i,h in enumerate(headers,3): ws.cell(5,i,h);ws.column_dimensions[chr(64+i)].width=34
        for r,row in enumerate(rows or [['No reviewed records available']],6):
            for c,value in enumerate(row,3): ws.cell(r,c,safe(value)).alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[r].height=55
        for cell in ws[5][2:]: cell.fill=PatternFill('solid',fgColor='274A45');cell.font=Font(color='FFFFFF',bold=True)
    for kind,name in [('sale','Sales Comps'),('rent','Rental Comps')]:
        table(name,['Address','Price','Price type','Event date','Status','Rationale','Source'],[[c['address'],float(c['price']),c['price_type'],c['event_date'] or 'Unknown',c['status'],c['rationale'],c['source_id']] for c in spec['comps'] if c['category']==kind])
    table('Property Evidence',['Claim','Type','Status','Sources','Limitations'],[[c['text'],c['kind'],c['status'],', '.join(c['source_ids']),c['limitation']] for c in spec['claims']])
    table('HOA Questions',['Question','Reason','Status'],[[q['question'],q['reason'],q['status']] for q in spec['questions']])
    table('Sources',['Title','Type','Retrieved','Event date','URL','Source ID'],[[s['title'],s['kind'],s['retrieved_at'],s['event_date'] or 'Unknown',s['url'] or 'Supplied / synthetic',s['id']] for s in spec['sources']])
    for i,note in enumerate(spec['finance']['notes'],27): wb['Summary'].cell(i,3,note)
    wb.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
    buf=io.BytesIO();wb.save(buf)
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    output=io.BytesIO()
    # Preserve formulas and seed reproducible cached values without altering inputs.
    with zipfile.ZipFile(buf) as source,zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                index=int(item.filename.split('sheet')[-1].split('.')[0])-1
                tree=ET.fromstring(data)
                for cell in tree.iter('{'+ns+'}c'):
                    key=(names[index],cell.get('r'))
                    if key not in caches: continue
                    value=caches[key]
                    v=cell.find('{'+ns+'}v')
                    if v is None: v=ET.SubElement(cell,'{'+ns+'}v')
                    if value is None: cell.set('t','str');v.text='n.a.'
                    else: cell.set('t','n');v.text=str(value)
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            target.writestr(item,data)
    path.write_bytes(output.getvalue())
    path.with_name('workbook-checks.json').write_text(json.dumps({'engine':'portable fallback','cache':'Python shared calculation graph','native_excel_recalculation':'not tested; enabled on open'},indent=2))
