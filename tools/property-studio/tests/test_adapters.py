from pathlib import Path
from property_studio.adapters import Cosmos,LanguageModel
from property_studio.config import Settings


def test_cosmos_sends_real_image_payload_and_auth(tmp_path,monkeypatch):
    monkeypatch.setenv('COSMOS3_REASON_URL','https://configured.example')
    monkeypatch.setenv('GPU_BEARER_TOKEN','test-only-token')
    monkeypatch.delenv('COSMOS3_REASON_MODEL',raising=False)
    calls=[]
    def request(method,url,**kwargs):
        calls.append((method,url,kwargs))
        if method=='GET': return {'data':[{'id':'discovered-cosmos'}]}
        return {'choices':[{'message':{'content':'{"description":"Visible room","room":"living room","criteria":[],"limitations":[]}'}}]}
    monkeypatch.setattr('property_studio.adapters.request_json',request)
    photo=tmp_path/'photo.jpg';photo.write_bytes(b'image-fixture')
    output,model=Cosmos(Settings.load()).analyze(photo,['visible room'])
    assert model=='discovered-cosmos'
    assert calls[-1][1]=='https://configured.example/v1/chat/completions'
    body=calls[-1][2]
    assert body['headers']['Authorization']=='Bearer test-only-token'
    assert body['json']['messages'][0]['content'][1]['image_url']['url'].startswith('data:image/jpeg;base64,')
    response_format=body['json']['response_format']
    assert response_format['type']=='json_schema'
    assert set(response_format['json_schema']['schema']['required'])=={'description','room','criteria','limitations'}


def test_cosmos_rejects_incomplete_observations(tmp_path,monkeypatch):
    import pytest
    from property_studio.adapters import ProviderError
    monkeypatch.setenv('COSMOS3_REASON_URL','https://configured.example')
    monkeypatch.setenv('COSMOS3_REASON_MODEL','configured-model')
    photo=tmp_path/'photo.jpg';photo.write_bytes(b'image-fixture')
    for content,reason,message in [
        ('{"description":"A room"}','stop','schema'),
        ('{"description":','length','output budget'),
        ('{"description":"A room","room":"room","criteria":[{"criterion":"condition"}],"limitations":[]}','stop','schema'),
    ]:
        monkeypatch.setattr('property_studio.adapters.request_json',lambda *a,**k:{'choices':[{'finish_reason':reason,'message':{'content':content}}]})
        with pytest.raises(ProviderError,match=message): Cosmos(Settings.load()).analyze(photo,[])


def test_wandb_uses_explicit_model_and_project(monkeypatch):
    monkeypatch.setenv('WANDB_API_KEY','test-only-key')
    monkeypatch.setenv('PROPERTY_LLM_MODEL','confirmed-model')
    monkeypatch.setenv('WANDB_TEAM','test-team')
    monkeypatch.setenv('WANDB_PROJECT','test-project')
    def request(method,url,**kwargs):
        assert kwargs['json']['model']=='confirmed-model'
        assert kwargs['headers']['OpenAI-Project']=='test-team/test-project'
        return {'choices':[{'message':{'content':'{"answer":"test"}'}}]}
    monkeypatch.setattr('property_studio.adapters.request_json',request)
    assert LanguageModel(Settings.load()).complete('instructions',{})=={'answer':'test'}


def test_portable_export_when_node_unavailable(tmp_path,monkeypatch):
    from openpyxl import load_workbook
    from property_studio.models import Run,Property
    from property_studio.service import synthetic_input
    from property_studio.finance import calculate
    from property_studio.workbook import export_workbook
    s=Settings.load();s.node_bin='/nonexistent-property-test-node'
    request=synthetic_input()
    run=Run(request=request,property=Property(address=request.address),finance=calculate(request.assumptions))
    export_workbook(run,tmp_path,s)
    cached=load_workbook(tmp_path/'analysis.xlsx',data_only=True)
    formulas=load_workbook(tmp_path/'analysis.xlsx',data_only=False)
    assert cached['Summary']['D10'].value==run.finance['metrics']['noi']
    assert formulas['Summary']['D10'].value.startswith('=')
    assert 'Sensitivity' in cached.sheetnames
